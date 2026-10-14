#!/usr/bin/env python3
"""Write the HIT140 Assessment 3 Word report (simple English, figures included).

    python code/clean_data.py
    python code/analytics.py
    python docs/generate_docx.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FIGS = HERE / "figures"
OUT = HERE / "HIT140_Assessment3_Group_Report.docx"
NAVY = RGBColor(0x1A, 0x1A, 0x1A)
BLUE = RGBColor(0x1F, 0x4E, 0x79)


def set_font(run, size=12, bold=False, italic=False, color=NAVY):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "8FAADC")
        borders.append(el)
    tcPr.append(borders)


def cell_text(cell, text, size=10, bold=False, center=False, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=RGBColor(0xFF, 0xFF, 0xFF) if white else NAVY)
    border(cell)


def para(doc, text, size=12, bold=False, italic=False, after=8, align="justify"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    size = 14 if level == 1 else 13
    for run in p.runs:
        set_font(run, size=size, bold=True, color=BLUE)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, size=10, bold=True, center=True, white=True)
        shade(t.rows[0].cells[i], "1F4E79")
        t.rows[0].cells[i].width = Cm(widths[i])
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell_text(t.rows[r].cells[c], val, size=10, center=(c > 0))
            t.rows[r].cells[c].width = Cm(widths[c])
            if r % 2 == 0:
                shade(t.rows[r].cells[c], "D6E3F0")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return t


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    para(doc, "CHARLES DARWIN UNIVERSITY", size=11, bold=True, after=2, align="center")
    para(doc, "HIT140 Foundations of Data Science", size=16, bold=True, after=2, align="center")
    para(doc, "Assessment 3: Group Project Report", size=14, bold=True, after=2, align="center")
    para(
        doc,
        "Objective 2 — Two linear regression models of FIFA World Cup 2026 matches",
        size=12,
        italic=True,
        after=10,
        align="center",
    )

    para(doc, "Group members (same group as Assessment 2)", size=12, bold=True, after=4, align="left")
    table(
        doc,
        ["Name", "Student ID", "Contribution"],
        [
            ["Bhupraj Tamang", "", ""],
            ["", "", ""],
            ["", "", ""],
            ["", "", ""],
        ],
        [6.5, 4.5, 6.0],
    )
    para(
        doc,
        "Python code and data: https://github.com/BhuprajTmg/data-project/tree/cursor/assessment3-submission-0741",
        size=11,
        italic=True,
        after=4,
        align="left",
    )
    para(
        doc,
        "Run:  python code/clean_data.py    then    python code/analytics.py",
        size=11,
        italic=True,
        after=8,
        align="left",
    )

    heading(doc, "1. What this report is about")
    para(
        doc,
        "This report is for Assessment 3, Objective 2. We try to predict football match results from the 2026 FIFA World Cup before the game starts. The World Cup was in Canada, Mexico and the USA from 11 June to 19 July 2026. There were 48 teams and 104 official matches. Spain beat Argentina 1–0 after extra time in the final.",
    )
    para(
        doc,
        "We cannot use the regression models from Assessment 2. Those models used FIFA rank, squad money value, average age, past titles, host, rest days, same confederation and knockout. The lecturer asked us to build new models. Each model has four variables. Only two variables are used in both models.",
    )
    para(
        doc,
        "Model 2.1 predicts the goal difference of a match (team 1 goals minus team 2 goals). It has 104 rows, one for each match. Model 2.2 predicts how many goals one team scores in one match. It has 208 rows, two for each match. We only use facts known before kick-off. We do not use shots, possession, cards or minutes from that match. All cleaning and modelling is in Python. Excel only stores the original tables. We use an 80/20 train/test split, seed 42, and α = 0.05.",
    )

    heading(doc, "2. Data and cleaning")
    heading(doc, "2.1 Original data", 2)
    para(
        doc,
        "The original file is data/original/fifa_data_original.xlsx. It has three sheets. Matches has 104 games. Teams has 48 countries. Players has 1,016 players. We used the player sheet only to check the data: 294 player goals plus 14 own goals equal 308 official goals, and there are 15 red cards. Player goals and cards from this World Cup are not put into the models.",
    )
    heading(doc, "2.2 Cleaning", 2)
    para(
        doc,
        "Cleaning is in code/clean_data.py. We make country names the same so the tables join (for example United States becomes USA). We fix dates and numbers. We drop duplicate rows. Blank group names in knockout games, and blank penalty winners, are filled with “Not applicable”. After this there are no missing values. We also make a roster table from player age and position only, not from minutes or goals in this tournament.",
    )
    heading(doc, "2.3 The four new variables", 2)
    para(
        doc,
        "Shared (2): how many World Cups the team played before, and how many days into the tournament the match is (from the published dates). Unique to Model 2.1: the gap in the share of players under 23, and the gap in how spread out the ages are. Unique to Model 2.2: how many World Cups the opponent played before, and whether the team is from UEFA. We do not use rank, value, average age, titles, host, rest, same confederation or knockout.",
    )

    heading(doc, "3. Linear Regression 2.1 — goal difference")
    para(
        doc,
        "The target is goal_diff. We have 104 matches: 83 for training and 21 for testing. Training R² = 0.321 (adjusted R² = 0.286). On new matches, test R² = 0.339, RMSE = 1.53 goals and MAE = 1.10 goals. Training and test R² are close, so the model is not only memorising the training matches.",
    )
    para(
        doc,
        "Table 1. Model 2.1 results. A star means p < 0.05.",
        size=11,
        italic=True,
        after=4,
        align="left",
    )
    table(
        doc,
        ["Variable", "Coefficient", "p-value", "Plain meaning"],
        [
            ["Intercept", "+0.489", "0.200", "Not significant"],
            ["appearances_diff *", "+0.1249", "0.000", "More past World Cups than the other team → more net goals"],
            ["days_into_tournament", "−0.0271", "0.232", "Match day; not significant"],
            ["youth_share_diff", "+2.798", "0.111", "More under-23 players; not significant"],
            ["age_spread_diff", "−0.175", "0.496", "Wider age mix; not significant"],
        ],
        [4.2, 2.4, 2.2, 8.2],
    )
    para(
        doc,
        "Only appearances_diff is significant. If team 1 has one extra past World Cup compared with team 2, we expect about +0.12 extra net goals. Youth share and age spread tell us about the mix of the squad, not the old average-age variable. They are not significant after we include World Cup experience. We still keep all four variables because we chose them before we looked at the p-values.",
    )

    heading(doc, "4. Linear Regression 2.2 — goals scored by one team")
    para(
        doc,
        "The target is goals scored by that team. We have 208 rows: 166 for training and 42 for testing. Training R² = 0.213 (adjusted R² = 0.193). Test R² = 0.074, RMSE = 1.34 and MAE = 1.04. This test fit is weak. Counting one team’s goals is harder than predicting a match margin, and we cannot bring back rank or value from Assessment 2.",
    )
    para(
        doc,
        "Table 2. Model 2.2 results. A star means p < 0.05.",
        size=11,
        italic=True,
        after=4,
        align="left",
    )
    table(
        doc,
        ["Variable", "Coefficient", "p-value", "Plain meaning"],
        [
            ["Intercept", "+1.342", "0.000", "Starting point when other X = 0"],
            ["team_appearances *", "+0.0553", "0.002", "More own World Cup history → more goals"],
            ["days_into_tournament", "+0.0059", "0.634", "Match day; not significant"],
            ["opp_appearances *", "−0.0654", "0.000", "Opponent with more World Cups → fewer goals"],
            ["is_uefa *", "+0.504", "0.023", "A UEFA team scores about +0.50 extra goals"],
        ],
        [4.2, 2.4, 2.2, 8.2],
    )
    para(
        doc,
        "Three results are significant: team_appearances, opp_appearances and is_uefa. A team with more World Cup history scores more. A more experienced opponent is linked to about 0.07 fewer goals per extra appearance. UEFA teams score about half a goal more after we hold experience fixed. The calendar day is not significant in either model.",
    )

    heading(doc, "5. Figures")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run()
    r.add_picture(str(FIGS / "reg1_pred.png"), width=Inches(3.05))
    p.add_run("  ")
    r2 = p.add_run()
    r2.add_picture(str(FIGS / "reg2_pred.png"), width=Inches(3.05))
    para(
        doc,
        "Figure 1 (left): Model 2.1 predicted vs actual goal difference. Figure 2 (right): Model 2.2 predicted vs actual goals. The dashed line is a perfect prediction. Points closer to the line are better predictions. These charts were made in Python by code/analytics.py.",
        size=10,
        italic=True,
        after=8,
        align="center",
    )

    heading(doc, "6. Checks, limits and conclusion")
    heading(doc, "6.1 Checks", 2)
    para(
        doc,
        "VIF is 1.21 in Model 2.1 and 1.25 in Model 2.2. Both are far below 5, so the variables are not too similar to each other. Shapiro–Wilk p = 0.802 for Model 2.1 (residuals look normal) and p = 0.0006 for Model 2.2 (residuals are not normal). Goals are 0, 1, 2, 3… so linear regression is a simple tool. The brief asked for linear regression, so we used it. A Poisson model would be a fair next step for Model 2.2.",
    )
    heading(doc, "6.2 Limits", 2)
    para(
        doc,
        "First, the 208 rows in Model 2.2 are not fully independent, because two rows come from the same match. Second, the test set in Model 2.1 is only 21 matches, so test R² = 0.339 is a rough check. Third, youth share and age spread come from the named squad, not from how many minutes those players later played. Fourth, we chose the four variables first. Dropping the weak ones after seeing p-values would be cheating the test.",
    )
    heading(doc, "6.3 Conclusion", 2)
    para(
        doc,
        "With new variables and no Assessment 2 predictors, past World Cup experience is the clearest signal. In Model 2.1, the appearance gap is the only significant driver of goal difference. In Model 2.2, a team scores more if it has more World Cup history, scores less against a more experienced opponent, and UEFA teams score about half a goal more. The goal-difference model predicts new matches much better (test R² 0.34 versus 0.07).",
        after=4,
    )

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
