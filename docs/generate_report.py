#!/usr/bin/env python3
"""Build the HIT140 Assessment 3 Word report.

Assignment format: A4, Times New Roman 12 pt, single-spaced,
numbered headings. Run from the repository root after analytics:

    python code/clean_data.py
    python code/analytics.py
    python docs/generate_report.py
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FIGS = Path(__file__).resolve().parent / "figures"
OUT = Path(__file__).resolve().parent / "HIT140_Assessment3_Group_Report.docx"

NAVY = RGBColor(0x1A, 0x1A, 0x1A)
HEADING = RGBColor(0x1F, 0x4E, 0x79)


def set_run_font(run, size=12, bold=False, italic=False, color=NAVY, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "8FAADC")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, size=10, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    set_cell_border(cell)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = NAVY
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_after = Pt(8)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")

    for style_name, size, space_before, space_after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 13, 10, 4),
    ):
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = HEADING
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.line_spacing = 1.0
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
        rfonts.set(qn("w:eastAsia"), "Times New Roman")


def add_para(doc, text, size=12, bold=False, italic=False, space_after=8, align="justify"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        set_run_font(run, size=14, bold=True, color=HEADING)
    return p


def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_run_font(run, size=13, bold=True, color=HEADING)
    return p


def make_table(doc, headers, rows, col_widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, size=10, bold=True, center=True)
        shade(table.rows[0].cells[i], "1F4E79")
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.bold = True
        table.rows[0].cells[i].width = Cm(col_widths[i])
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell_text(table.rows[r].cells[c], val, size=10, center=(c > 0))
            table.rows[r].cells[c].width = Cm(col_widths[c])
            if r % 2 == 0:
                shade(table.rows[r].cells[c], "D6E3F0")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def build():
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    add_para(doc, "CHARLES DARWIN UNIVERSITY", size=11, bold=True, space_after=2, align="center")
    add_para(doc, "HIT140 Foundations of Data Science", size=16, bold=True, space_after=2, align="center")
    add_para(doc, "Assessment 3: Group Project Report", size=14, bold=True, space_after=2, align="center")
    add_para(
        doc,
        "Objective 2 — Linear regression models of FIFA World Cup 2026 match outcomes",
        size=12,
        italic=True,
        space_after=10,
        align="center",
    )

    add_para(doc, "Group members (same group as Assessment 2)", size=12, bold=True, space_after=4, align="left")
    make_table(
        doc,
        ["Name", "Student ID", "Contribution"],
        [
            ["Bhupraj Tamang", "", ""],
            ["", "", ""],
            ["", "", ""],
            ["", "", ""],
        ],
        [6.5, 4.5, 6.0],
    )
    add_para(
        doc,
        "Python code and dataset: https://github.com/BhuprajTmg/data-project/tree/cursor/assessment3-submission-0741",
        size=11,
        italic=True,
        space_after=4,
        align="left",
    )
    add_para(
        doc,
        "Run order:  python code/clean_data.py    then    python code/analytics.py",
        size=11,
        italic=True,
        space_after=8,
        align="left",
    )

    add_h1(doc, "1. Introduction")
    add_para(
        doc,
        "This report answers Objective 2 of the HIT140 project brief. The task is to predict FIFA World Cup 2026 match outcomes using only information that is known before kick-off. The tournament was hosted by Canada, Mexico and the United States from 11 June to 19 July 2026. Forty-eight teams played 104 official matches. Spain defeated Argentina 1–0 after extra time in the final.",
    )
    add_para(
        doc,
        "We fit two ordinary least squares (OLS) linear regression models. Model 2.1 predicts the goal difference of a match (team 1 goals minus team 2 goals) using 104 rows and eight explanatory variables. Model 2.2 predicts the number of goals scored by one team in one match using 208 rows (two rows per match) and eight explanatory variables. The brief allows at most four of those variables to be shared. Every predictor is known before the referee starts the game. Shots, possession, cards in that match, and other in-match events are excluded because they would leak the result we are trying to predict.",
    )
    add_para(
        doc,
        "All wrangling, modelling and figures were produced in Python. Microsoft Excel was used only to store the original raw tables. Both models use an 80/20 training/test split, random seed 42, and significance level α = 0.05.",
    )

    add_h1(doc, "2. Data and cleaning")
    add_h2(doc, "2.1 Original data")
    add_para(
        doc,
        "The original workbook is data/original/fifa_data_original.xlsx. It contains three sheets compiled from the FIFA match centre, Wikipedia, Transfermarkt and the FIFA ranking freeze of 11 June 2026. The matches sheet has 104 games (date, stage, teams, scores, attendance). The teams sheet has 48 national sides (ranking, confederation, host flag, previous titles, previous World Cup appearances, squad value in EUR million, average age). The players sheet has 1,016 players who featured. Players are not used as regression predictors. They are a quality check: 294 player goals plus 14 own goals equal the official 308-goal tournament total, and 15 red cards match the official count.",
    )
    add_h2(doc, "2.2 Cleaning steps")
    add_para(
        doc,
        "Cleaning is in code/clean_data.py. Country names are mapped to one spelling so the three tables join (for example United States → USA, Turkey → Turkiye, Czech Republic → Czechia). Dates are parsed. Numeric fields are coerced. Duplicate match numbers, team names and player–team rows are dropped. The host flag is stored as 0/1. Knockout matches have no group, so the 32 blank group cells are labelled “Not applicable”. Penalty winner is blank unless the match went to penalties; those blanks are also labelled “Not applicable”. After cleaning there are no missing values in the three tables. The cleaned files are saved in data/cleaned/.",
    )
    add_h2(doc, "2.3 Feature construction for the two models")
    add_para(
        doc,
        "A knockout flag is 1 if the stage is Round of 32 through the Final (32 matches) and 0 for the group stage (72 matches). Rest days are the number of days since that team’s previous tournament match. A first match has no previous date, so those 48 opening games are filled with the median rest from the rest of the tournament (5 days). That fill uses the published schedule, not the score.",
    )
    add_para(
        doc,
        "Shared variables (exactly four): knockout, FIFA rank, squad market value, and rest days. Unique to Model 2.1: age_diff, titles_diff, host_diff, same_confed. Unique to Model 2.2: team_appearances, opp_appearances, is_uefa, days_into_tournament (days since 11 June, from the calendar). Model 2.1 stores rank, value and rest as team-1-minus-team-2 differences. Model 2.2 stores them as the team’s own rank, value and rest. A smaller FIFA rank number means a stronger team.",
    )

    add_h1(doc, "3. Linear Regression 2.1 — goal difference")
    add_para(
        doc,
        "The target is goal_diff. The sample is 104 matches, split into 83 training matches and 21 test matches. Training R² = 0.493 (adjusted R² = 0.438). On the held-out matches, test R² = 0.306, RMSE = 1.57 goals and MAE = 1.18 goals. The model describes about half of the training variation. Test fit is lower, which is expected with 21 held-out matches, but the predictions are still better than using the mean.",
    )
    add_para(doc, "Table 1. OLS coefficients for Model 2.1 (training sample). An asterisk marks p < 0.05.", size=11, italic=True, space_after=4, align="left")
    make_table(
        doc,
        ["Variable", "Coefficient", "p-value", "Interpretation"],
        [
            ["Intercept", "+0.184", "0.383", "Not significant"],
            ["rank_diff *", "−0.0232", "0.002", "Better FIFA rank for team 1 → more net goals"],
            ["value_diff *", "+0.0017", "0.004", "EUR 1 billion extra squad value ≈ +1.7 net goals"],
            ["rest_diff", "+0.172", "0.641", "Not significant"],
            ["knockout", "−0.726", "0.097", "Knockouts look tighter (marginal only)"],
            ["age_diff", "+0.113", "0.427", "Not significant (unique to 2.1)"],
            ["titles_diff", "−0.132", "0.336", "Not significant (unique to 2.1)"],
            ["host_diff", "+0.366", "0.494", "Not significant (unique to 2.1)"],
            ["same_confed", "+0.218", "0.684", "Not significant (unique to 2.1)"],
        ],
        [3.4, 2.6, 2.2, 8.8],
    )
    add_para(
        doc,
        "Only rank_diff and value_diff are significant at 5%. The rank coefficient is negative because rank 1 is better than rank 20. If team 1 is stronger, rank_diff is negative; a negative coefficient then predicts a positive goal difference for team 1. Squad value has the expected sign: the richer squad is predicted to outscore the poorer one. Age, titles, host and same confederation add little once rank and value are already in the model. They are kept because the brief asked for eight thoughtful variables, not a search for small p-values.",
    )

    add_h1(doc, "4. Linear Regression 2.2 — goals scored by one team")
    add_para(
        doc,
        "The target is goals scored by that team in that match. The sample is 208 rows, split into 166 training rows and 42 test rows. Training R² = 0.264 (adjusted R² = 0.227). Test R² = 0.074, RMSE = 1.34 and MAE = 1.01. Test fit is weak. Two reasons explain this. First, a single team’s goal count is noisier than a match margin. Second, the sharing rule stops opponent rank and opponent value from appearing in this model. The weak test R² is reported rather than breaking the four-shared-variable limit.",
    )
    add_para(doc, "Table 2. OLS coefficients for Model 2.2 (training sample). An asterisk marks p < 0.05.", size=11, italic=True, space_after=4, align="left")
    make_table(
        doc,
        ["Variable", "Coefficient", "p-value", "Interpretation"],
        [
            ["Intercept", "+3.14", "0.000", "Baseline when other X = 0"],
            ["team_rank *", "−0.0133", "0.037", "Better own rank → more goals"],
            ["team_value", "+0.0003", "0.535", "Not significant (shared idea: value)"],
            ["rest_days", "−0.207", "0.102", "Not significant (shared idea: rest)"],
            ["knockout", "−0.372", "0.313", "Not significant (shared idea: knockout)"],
            ["team_appearances", "+0.015", "0.558", "Own prior World Cup appearances"],
            ["opp_appearances *", "−0.0645", "0.000", "More World Cup history for the opponent → fewer goals"],
            ["is_uefa", "+0.303", "0.241", "UEFA dummy, not significant"],
            ["days_into_tournament", "+0.008", "0.668", "Schedule day, not significant"],
        ],
        [3.4, 2.6, 2.2, 8.8],
    )
    add_para(
        doc,
        "Two coefficients are significant: team_rank and opp_appearances. A better-ranked team scores more. Facing an opponent with more previous World Cups is linked to about 0.06 fewer goals per extra appearance. Value, rest, knockout, UEFA membership and calendar day are not significant after those two variables are included.",
    )

    add_h1(doc, "5. Visualisation")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run()
    r.add_picture(str(FIGS / "reg1_pred.png"), width=Inches(3.05))
    p.add_run("  ")
    r2 = p.add_run()
    r2.add_picture(str(FIGS / "reg2_pred.png"), width=Inches(3.05))
    add_para(
        doc,
        "Figure 1 (left) and Figure 2 (right). Test-set predicted versus actual values. Left: Model 2.1 goal difference. Right: Model 2.2 team goals. The dashed line is a perfect prediction. Both charts were produced by code/analytics.py.",
        size=10,
        italic=True,
        space_after=8,
        align="center",
    )

    add_h1(doc, "6. Diagnostics, limitations and conclusion")
    add_h2(doc, "6.1 Diagnostics")
    add_para(
        doc,
        "The largest variance inflation factor, excluding the intercept, is 3.30 in Model 2.1 and 3.44 in Model 2.2. Both are below 5, so multicollinearity is not a problem. Shapiro–Wilk tests on the training residuals give p = 0.937 for Model 2.1 (consistent with normal residuals) and p = 0.017 for Model 2.2 (residuals are not normal). Goals are small counts (0, 1, 2, 3…), so OLS is a blunt tool. The brief asked for linear regression, which is why OLS is used. A Poisson model would be a reasonable next step for Model 2.2.",
    )
    add_h2(doc, "6.2 Limitations")
    add_para(
        doc,
        "Five limitations should be kept in mind. First, the 208 rows in Model 2.2 are not independent, because two rows come from the same match; this can make standard errors look too small. Model 2.1 does not have that problem. Second, Transfermarkt squad values and average ages are a snapshot and can move with the pull date; rankings, scores and titles do not. Third, the test set for Model 2.1 is only 21 matches, so test R² = 0.306 is a rough check. Fourth, first-match rest is a median fill. Fifth, the eight variables were chosen in advance. Dropping the non-significant ones after seeing p-values would be p-hacking.",
    )
    add_h2(doc, "6.3 Conclusion")
    add_para(
        doc,
        "With only pre-match information, FIFA ranking is the most stable signal. In Model 2.1, the rank gap and the squad-value gap are the significant drivers of goal difference. In Model 2.2, own rank and the opponent’s World Cup history are the significant drivers of goals scored. The goal-difference model predicts held-out matches much better (test R² 0.31 versus 0.07). That gap is partly because a match margin is easier than a count, and partly because the four-variable sharing rule blocks opponent rank from appearing in both models.",
        space_after=4,
    )

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
